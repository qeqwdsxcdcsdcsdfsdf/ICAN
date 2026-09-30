from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_project_doc():
    doc = Document()
    
    style = doc.styles['Heading 1']
    style.font.size = Pt(20)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 51, 102)
    
    style2 = doc.styles['Heading 2']
    style2.font.size = Pt(16)
    style2.font.bold = True
    style2.font.color.rgb = RGBColor(0, 51, 102)
    
    style3 = doc.styles['Heading 3']
    style3.font.size = Pt(14)
    style3.font.bold = True
    style3.font.color.rgb = RGBColor(0, 102, 204)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("人脸考勤GPS定位系统\n项目文档")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    doc.add_heading('1. 项目概述', level=1)
    doc.add_paragraph('本项目是一个基于人脸识别和GPS定位的考勤管理系统，主要用于企业或学校的员工/学生考勤管理。系统包含三个主要模块：Android移动端App、Spring Boot后端服务和PC端管理员管理系统。')
    
    doc.add_heading('2. 技术架构', level=1)
    doc.add_heading('2.1 整体架构', level=2)
    doc.add_paragraph('系统采用三层架构设计：')
    doc.add_paragraph('• 前端层：Android App（员工使用）和Vue.js PC管理端（管理员使用）')
    doc.add_paragraph('• 服务层：Spring Boot后端服务，处理业务逻辑和数据访问')
    doc.add_paragraph('• 数据层：MySQL数据库，存储员工信息、考勤记录、定位配置等')
    
    doc.add_heading('2.2 技术栈', level=2)
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '模块'
    hdr_cells[1].text = '技术栈'
    row_cells = table.rows[1].cells
    row_cells[0].text = 'Android App'
    row_cells[1].text = 'Java, Retrofit, OkHttp, FaceDetector'
    row_cells = table.rows[2].cells
    row_cells[0].text = '后端服务'
    row_cells[1].text = 'Spring Boot, MyBatis-Plus, MySQL, PCA人脸识别'
    row_cells = table.rows[3].cells
    row_cells[0].text = 'PC管理端'
    row_cells[1].text = 'Vue.js, Element UI, Axios'
    
    doc.add_heading('3. 功能模块', level=1)
    doc.add_heading('3.1 Android App功能', level=2)
    doc.add_paragraph('• 用户登录：员工使用工号和密码登录')
    doc.add_paragraph('• 人脸注册：员工注册人脸信息')
    doc.add_paragraph('• 人脸签到：通过人脸识别进行考勤签到')
    doc.add_paragraph('• GPS定位：获取当前位置，计算与考勤点的距离')
    doc.add_paragraph('• 考勤记录：查看个人考勤历史记录')
    doc.add_paragraph('• 修改密码：员工修改个人密码')
    
    doc.add_heading('3.2 后端服务功能', level=2)
    doc.add_paragraph('• 用户认证：员工登录验证')
    doc.add_paragraph('• 人脸注册：存储人脸特征信息')
    doc.add_paragraph('• 人脸验证：PCA算法进行人脸比对')
    doc.add_paragraph('• 签到管理：处理签到请求，验证位置和人脸')
    doc.add_paragraph('• 定位配置：管理考勤点位置和半径')
    doc.add_paragraph('• 考勤记录：查询和管理考勤记录')
    
    doc.add_heading('3.3 PC管理端功能', level=2)
    doc.add_paragraph('• 管理员登录')
    doc.add_paragraph('• 员工管理：添加、编辑、删除员工信息')
    doc.add_paragraph('• 考勤记录：查看、筛选、删除考勤记录')
    doc.add_paragraph('• 定位配置：设置考勤中心点和半径')
    doc.add_paragraph('• 反向地理编码：根据坐标获取地址名称')
    
    doc.add_heading('4. 项目结构', level=1)
    doc.add_paragraph('FaceAttendanceGPSTotal/')
    doc.add_paragraph('├── 01-Android-App/          # Android移动端App')
    doc.add_paragraph('│   ├── app/src/main/java/com/attendance/')
    doc.add_paragraph('│   │   ├── activity/        # 活动页面')
    doc.add_paragraph('│   │   ├── api/             # API接口')
    doc.add_paragraph('│   │   ├── bean/            # 数据模型')
    doc.add_paragraph('│   │   └── util/            # 工具类')
    doc.add_paragraph('│   └── app/src/main/res/    # 资源文件')
    doc.add_paragraph('├── 02-SpringBoot-Server/    # Spring Boot后端服务')
    doc.add_paragraph('│   ├── src/main/java/com/attendance/')
    doc.add_paragraph('│   │   ├── controller/      # 控制器')
    doc.add_paragraph('│   │   ├── service/         # 服务层')
    doc.add_paragraph('│   │   ├── mapper/          # 数据访问层')
    doc.add_paragraph('│   │   ├── entity/          # 实体类')
    doc.add_paragraph('│   │   ├── config/          # 配置类')
    doc.add_paragraph('│   │   └── util/            # 工具类')
    doc.add_paragraph('│   └── src/main/resources/  # 配置文件')
    doc.add_paragraph('└── 03-PC-Admin-Vue/         # PC端管理员系统')
    doc.add_paragraph('    ├── src/views/           # 页面组件')
    doc.add_paragraph('    ├── src/api/             # API接口')
    doc.add_paragraph('    └── src/components/      # 公共组件')
    
    doc.save('C:\\Users\\35416\\Desktop\\FaceAttendanceGPSTotal\\项目文档.docx')
    print('项目文档.docx 已生成')

def create_run_doc():
    doc = Document()
    
    style = doc.styles['Heading 1']
    style.font.size = Pt(20)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 51, 102)
    
    style2 = doc.styles['Heading 2']
    style2.font.size = Pt(16)
    style2.font.bold = True
    style2.font.color.rgb = RGBColor(0, 51, 102)
    
    style3 = doc.styles['Heading 3']
    style3.font.size = Pt(14)
    style3.font.bold = True
    style3.font.color.rgb = RGBColor(0, 102, 204)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("人脸考勤GPS定位系统\n运行文档")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    doc.add_heading('1. 环境要求', level=1)
    doc.add_paragraph('• JDK 1.8 或更高版本')
    doc.add_paragraph('• MySQL 5.7 或更高版本')
    doc.add_paragraph('• Android Studio（用于构建App）')
    doc.add_paragraph('• Node.js 14+（用于构建PC管理端）')
    doc.add_paragraph('• Maven（后端构建工具）')
    
    doc.add_heading('2. 数据库配置', level=1)
    doc.add_paragraph('2.1 创建数据库')
    doc.add_paragraph('打开MySQL客户端，执行以下命令：')
    p = doc.add_paragraph()
    run = p.add_run('CREATE DATABASE face_attendance_gps CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_paragraph('2.2 导入初始化数据')
    doc.add_paragraph('执行 02-SpringBoot-Server/src/main/resources/init.sql 文件')
    
    doc.add_paragraph('2.3 修改数据库连接配置')
    doc.add_paragraph('编辑 02-SpringBoot-Server/src/main/resources/application.yml 文件：')
    p = doc.add_paragraph()
    run = p.add_run('spring:\n  datasource:\n    url: jdbc:mysql://localhost:3306/face_attendance_gps?useUnicode=true&characterEncoding=utf8&serverTimezone=Asia/Shanghai\n    username: root\n    password: 你的密码')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_heading('3. 启动后端服务', level=1)
    doc.add_paragraph('方法一：使用 start.bat 脚本（推荐）')
    doc.add_paragraph('双击运行 FaceAttendanceGPSTotal/start.bat')
    
    doc.add_paragraph('方法二：使用Maven命令')
    p = doc.add_paragraph()
    run = p.add_run('cd 02-SpringBoot-Server\nmvnw.cmd clean package -DskipTests\njava -jar target/face-attendance-gps-1.0.0.jar')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_paragraph('服务启动后访问地址：http://localhost:8080')
    
    doc.add_heading('4. 启动PC管理端', level=1)
    doc.add_paragraph('方法一：使用已构建的静态文件')
    doc.add_paragraph('将 03-PC-Admin-Vue/dist/ 目录部署到Web服务器（如Nginx）')
    
    doc.add_paragraph('方法二：开发模式运行')
    p = doc.add_paragraph()
    run = p.add_run('cd 03-PC-Admin-Vue\nnpm install\nnpm run dev')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_paragraph('默认管理员账号：admin / 123456')
    
    doc.add_heading('5. 构建和安装Android App', level=1)
    doc.add_paragraph('5.1 构建APK')
    p = doc.add_paragraph()
    run = p.add_run('cd 01-Android-App\ngradlew.bat assembleDebug')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_paragraph('5.2 获取APK文件')
    doc.add_paragraph('构建成功后，APK文件位于：')
    p = doc.add_paragraph()
    run = p.add_run('01-Android-App/app/build/outputs/apk/debug/app-debug.apk')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_paragraph('5.3 安装到手机')
    doc.add_paragraph('• 将APK文件复制到Android手机')
    doc.add_paragraph('• 在手机上打开文件管理器，找到APK文件')
    doc.add_paragraph('• 点击安装（如果提示"未知来源"，请允许安装）')
    
    doc.add_heading('6. 测试账号', level=1)
    table = doc.add_table(rows=3, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '角色'
    hdr_cells[1].text = '账号'
    hdr_cells[2].text = '密码'
    row_cells = table.rows[1].cells
    row_cells[0].text = '管理员'
    row_cells[1].text = 'admin'
    row_cells[2].text = '123456'
    row_cells = table.rows[2].cells
    row_cells[0].text = '员工'
    row_cells[1].text = 's004'
    row_cells[2].text = '123456'
    
    doc.save('C:\\Users\\35416\\Desktop\\FaceAttendanceGPSTotal\\运行文档.docx')
    print('运行文档.docx 已生成')

def create_code_doc():
    doc = Document()
    
    style = doc.styles['Heading 1']
    style.font.size = Pt(20)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 51, 102)
    
    style2 = doc.styles['Heading 2']
    style2.font.size = Pt(16)
    style2.font.bold = True
    style2.font.color.rgb = RGBColor(0, 51, 102)
    
    style3 = doc.styles['Heading 3']
    style3.font.size = Pt(14)
    style3.font.bold = True
    style3.font.color.rgb = RGBColor(0, 102, 204)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("人脸考勤GPS定位系统\n代码说明文档")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    doc.add_heading('1. Android App代码说明', level=1)
    
    doc.add_heading('1.1 核心Activity', level=2)
    doc.add_paragraph('• LoginActivity.java - 用户登录页面')
    doc.add_paragraph('• MainActivity.java - 主页面，包含底部导航')
    doc.add_paragraph('• SignInActivity.java - 签到页面')
    doc.add_paragraph('• SignInFragment.java - 签到Fragment（嵌入MainActivity）')
    doc.add_paragraph('• FaceRecognitionActivity.java - 人脸识别页面')
    doc.add_paragraph('• FaceRegisterActivity.java - 人脸注册页面')
    doc.add_paragraph('• ChangePasswordActivity.java - 修改密码页面')
    doc.add_paragraph('• AttendanceRecordActivity.java - 考勤记录页面')
    
    doc.add_heading('1.2 API接口层', level=2)
    doc.add_paragraph('• ApiService.java - Retrofit配置和服务管理')
    doc.add_paragraph('  - BASE_URL：后端服务地址，默认 http://192.168.2.113:8080')
    doc.add_paragraph('  - getAttendanceApi()：获取API接口实例')
    doc.add_paragraph('• AttendanceApi.java - API接口定义')
    doc.add_paragraph('  - login()：登录接口')
    doc.add_paragraph('  - signIn()：签到接口（使用@Body传递JSON参数）')
    doc.add_paragraph('  - registerFace()：人脸注册接口')
    doc.add_paragraph('  - verifyFace()：人脸验证接口')
    doc.add_paragraph('  - getAttendanceRecords()：获取考勤记录')
    
    doc.add_heading('1.3 工具类', level=2)
    doc.add_paragraph('• LocationHelper.java - GPS定位辅助类')
    doc.add_paragraph('• FaceDetector.java - 人脸检测工具')
    doc.add_paragraph('• PermissionUtil.java - 权限申请工具')
    
    doc.add_heading('2. 后端服务代码说明', level=1)
    
    doc.add_heading('2.1 控制器层', level=2)
    doc.add_paragraph('• StaffController.java - 员工管理控制器')
    doc.add_paragraph('  - /api/staff/login - 员工登录')
    doc.add_paragraph('  - /api/staff/info - 获取员工信息')
    doc.add_paragraph('  - /api/staff/list - 获取员工列表')
    doc.add_paragraph('  - /api/staff/add - 添加员工')
    doc.add_paragraph('  - /api/staff/update - 更新员工')
    doc.add_paragraph('  - /api/staff/delete - 删除员工')
    doc.add_paragraph('• SignController.java - 签到管理控制器')
    doc.add_paragraph('  - /api/sign/in - 签到接口')
    doc.add_paragraph('  - /api/sign/verify-face - 人脸验证')
    doc.add_paragraph('  - /api/sign/records - 查询考勤记录')
    doc.add_paragraph('  - /api/sign/delete - 删除考勤记录')
    doc.add_paragraph('  - /api/sign/reverse-geocode - 反向地理编码')
    doc.add_paragraph('• LocationConfigController.java - 定位配置控制器')
    doc.add_paragraph('  - /api/location/config - 获取定位配置')
    doc.add_paragraph('  - /api/location/update - 更新定位配置')
    doc.add_paragraph('  - /api/location/check - 检查位置是否在范围内')
    
    doc.add_heading('2.2 服务层', level=2)
    doc.add_paragraph('• SignService.java - 签到服务')
    doc.add_paragraph('  - signIn()：签到核心逻辑')
    doc.add_paragraph('  - verifyFace()：人脸验证')
    doc.add_paragraph('  - isDuplicateSign()：重复签到判断')
    doc.add_paragraph('• StaffService.java - 员工服务')
    doc.add_paragraph('• LocationCheckService.java - 位置检查服务')
    
    doc.add_heading('2.3 工具类', level=2)
    doc.add_paragraph('• PCAUtil.java - PCA人脸特征提取和比对')
    doc.add_paragraph('  - extractFeature()：从图片提取人脸特征')
    doc.add_paragraph('  - compare()：人脸比对，返回相似度')
    doc.add_paragraph('• GeoUtil.java - 地理计算工具')
    doc.add_paragraph('  - calculateDistance()：计算两点间距离')
    doc.add_paragraph('• FaceCompareUtil.java - 人脸比对工具')
    
    doc.add_heading('2.4 实体类', level=2)
    doc.add_paragraph('• Staff.java - 员工实体')
    doc.add_paragraph('• Attendance.java - 考勤记录实体')
    doc.add_paragraph('• LocationConfig.java - 定位配置实体')
    doc.add_paragraph('• Admin.java - 管理员实体')
    
    doc.add_heading('3. PC管理端代码说明', level=1)
    
    doc.add_heading('3.1 页面组件', level=2)
    doc.add_paragraph('• StaffList.vue - 员工列表页面')
    doc.add_paragraph('• SignRecords.vue - 考勤记录页面')
    doc.add_paragraph('• LocationSet.vue - 定位配置页面')
    doc.add_paragraph('• Login.vue - 登录页面')
    
    doc.add_heading('3.2 API接口', level=2)
    doc.add_paragraph('• api/index.js - 统一API接口管理')
    
    doc.save('C:\\Users\\35416\\Desktop\\FaceAttendanceGPSTotal\\代码说明文档.docx')
    print('代码说明文档.docx 已生成')

def create_ip_doc():
    doc = Document()
    
    style = doc.styles['Heading 1']
    style.font.size = Pt(20)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 51, 102)
    
    style2 = doc.styles['Heading 2']
    style2.font.size = Pt(16)
    style2.font.bold = True
    style2.font.color.rgb = RGBColor(0, 51, 102)
    
    style3 = doc.styles['Heading 3']
    style3.font.size = Pt(14)
    style3.font.bold = True
    style3.font.color.rgb = RGBColor(0, 102, 204)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("人脸考勤GPS定位系统\nIP地址修改说明文档")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    doc.add_heading('1. 为什么需要修改IP', level=1)
    doc.add_paragraph('系统默认配置的服务器地址是局域网IP（192.168.2.113），当服务器部署在不同网络环境或需要外网访问时，需要修改服务器地址配置。')
    
    doc.add_heading('2. 需要修改的位置', level=1)
    
    doc.add_heading('2.1 Android App端', level=2)
    doc.add_paragraph('文件路径：')
    p = doc.add_paragraph()
    run = p.add_run('01-Android-App/app/src/main/java/com/attendance/api/ApiService.java')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_paragraph('修改内容：')
    p = doc.add_paragraph()
    run = p.add_run('private static final String BASE_URL = "http://你的服务器IP:8080";')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_paragraph('修改方法：')
    doc.add_paragraph('1. 使用Android Studio打开项目')
    doc.add_paragraph('2. 找到 ApiService.java 文件')
    doc.add_paragraph('3. 修改 BASE_URL 常量为新的服务器地址')
    doc.add_paragraph('4. 重新构建APK')
    
    doc.add_paragraph('注意：也可以在App设置页面修改服务器地址，无需重新构建。')
    doc.add_paragraph('App内修改路径：设置 → 服务器地址')
    
    doc.add_heading('2.2 PC管理端', level=2)
    doc.add_paragraph('文件路径：')
    p = doc.add_paragraph()
    run = p.add_run('03-PC-Admin-Vue/src/api/index.js')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_paragraph('修改内容：')
    p = doc.add_paragraph()
    run = p.add_run('const baseURL = "http://你的服务器IP:8080";')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_paragraph('修改方法：')
    doc.add_paragraph('1. 使用VS Code打开项目')
    doc.add_paragraph('2. 找到 src/api/index.js 文件')
    doc.add_paragraph('3. 修改 baseURL 变量为新的服务器地址')
    doc.add_paragraph('4. 重新构建前端（npm run build）')
    doc.add_paragraph('5. 部署新的dist目录')
    
    doc.add_heading('2.3 前端代理配置（开发环境）', level=2)
    doc.add_paragraph('文件路径：')
    p = doc.add_paragraph()
    run = p.add_run('03-PC-Admin-Vue/vue.config.js')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_paragraph('修改内容（开发环境代理）：')
    p = doc.add_paragraph()
    run = p.add_run('proxy: {\n  "/api": {\n    target: "http://你的服务器IP:8080",\n    changeOrigin: true\n  }\n}')
    run.font.family = 'Consolas'
    run.font.size = Pt(10)
    
    doc.add_heading('3. 修改步骤总结', level=1)
    table = doc.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '序号'
    hdr_cells[1].text = '修改位置'
    hdr_cells[2].text = '修改内容'
    row_cells = table.rows[1].cells
    row_cells[0].text = '1'
    row_cells[1].text = 'ApiService.java'
    row_cells[2].text = 'BASE_URL常量'
    row_cells = table.rows[2].cells
    row_cells[0].text = '2'
    row_cells[1].text = 'api/index.js'
    row_cells[2].text = 'baseURL变量'
    row_cells = table.rows[3].cells
    row_cells[0].text = '3'
    row_cells[1].text = 'vue.config.js'
    row_cells[2].text = 'proxy.target'
    
    doc.add_heading('4. 验证修改', level=1)
    doc.add_paragraph('修改完成后，验证步骤：')
    doc.add_paragraph('1. 重启后端服务')
    doc.add_paragraph('2. 重新构建Android App并安装')
    doc.add_paragraph('3. 重新构建PC管理端并部署')
    doc.add_paragraph('4. 测试登录功能，确认能正常连接服务器')
    doc.add_paragraph('5. 测试签到功能，确认人脸识别和GPS定位正常工作')
    
    doc.add_heading('5. 注意事项', level=1)
    doc.add_paragraph('• 确保服务器防火墙已开放8080端口')
    doc.add_paragraph('• 如果使用公网IP，需要配置端口映射')
    doc.add_paragraph('• 修改IP后需要重新构建相关模块')
    doc.add_paragraph('• Android App也可以在设置页面动态修改服务器地址')
    
    doc.save('C:\\Users\\35416\\Desktop\\FaceAttendanceGPSTotal\\IP地址修改说明文档.docx')
    print('IP地址修改说明文档.docx 已生成')

if __name__ == '__main__':
    create_project_doc()
    create_run_doc()
    create_code_doc()
    create_ip_doc()
    print('\n所有文档已生成完成！')